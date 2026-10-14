"""
competency.py
--------------
Competency assessment for mentorship / support supervision visits.

Design
  * ONE checklist drives everything. It lives in the `CompetencyItems` tab of your Google
    Sheet. The in-app form reads it live, and the Kobo form (XLSForm) is GENERATED from it,
    so changing your tools later means editing the checklist, not code.
  * Each health worker is assessed separately. Results are stored "long": one row per
    item scored, per assessment, in the `Competency` tab.
  * Baseline vs follow-up is simply the assessment type; progress compares the two.

Things you will probably want to change later (all in this file, clearly marked):
  SCALE, COMPETENT_MIN, CADRES, and the sample checklist in DEFAULT_ITEMS.
"""

import io
import re
import uuid
from datetime import datetime

import pandas as pd

import gsheets

# ---------------------------------------------------------------------
# EDIT ME: scoring scale. (code stored in the Sheet, label shown to users, numeric value)
# Numeric None = not counted (e.g. Not applicable).
# ---------------------------------------------------------------------
SCALE = [
    ("2", "Competent", 2.0),
    ("1", "Needs improvement", 1.0),
    ("0", "Not competent", 0.0),
    ("NA", "Not applicable", None),
]
COMPETENT_MIN = 2.0          # a score at or above this counts as "competent"

# EDIT ME: cadres offered when recording who is being assessed.
CADRES = [
    "Medical Officer", "Clinical Officer", "Nurse", "Midwife", "Laboratory Technician",
    "Pharmacy / Dispenser", "Records / Data Clerk", "Community Health Worker / VHT", "Other",
]

ASSESSMENT_TYPES = [("baseline", "Baseline"), ("follow_up", "Follow-up")]

SCORE_LABEL = {c: l for c, l, _ in SCALE}
SCORE_CODE = {l: c for c, l, _ in SCALE}
SCORE_VALUE = {c: v for c, _, v in SCALE}
TYPE_LABEL = dict(ASSESSMENT_TYPES)
TYPE_CODE = {l: c for c, l in ASSESSMENT_TYPES}

ITEM_COLS = ["item_id", "service_area", "domain", "competency", "active"]
COMP_COLS = ["assessment_id", "visit_ref", "facility", "mentee_name", "cadre", "service_area",
             "assessment_type", "assessment_date", "item_id", "competency", "score", "assessor",
             "comments", "timestamp", "kobo_id"]

# Names used by the generated Kobo form for its own questions — items can't reuse them.
RESERVED_NAMES = {"facility", "mentee_name", "cadre", "service_area", "assessment_type",
                  "assessment_date", "assessor", "visit_ref", "comments", "meta", "instanceid"}


# ---------------------------------------------------------------------
# EDIT ME: SAMPLE checklist — placeholders so the app works on day one.
# Replace with your programme's real competencies (national guidelines / your tools).
# Not clinical guidance. Format: area -> [(domain, [(item_id, competency), ...]), ...]
# ---------------------------------------------------------------------
_SAMPLE = {
    "ANC": [
        ("History and examination", [
            ("anc_01", "Takes a focused obstetric history (LMP, gravida/para, danger signs)"),
            ("anc_02", "Measures blood pressure correctly and interprets the result"),
            ("anc_03", "Performs abdominal examination and fundal height measurement"),
        ]),
        ("Screening and prevention", [
            ("anc_04", "Offers HIV testing and syphilis screening to every client"),
            ("anc_05", "Gives IPTp-SP and iron/folic acid supplements correctly"),
            ("anc_06", "Identifies danger signs and refers appropriately"),
        ]),
        ("Counselling and records", [
            ("anc_07", "Counsels on birth preparedness and danger signs"),
            ("anc_08", "Completes the ANC register and card accurately"),
        ]),
    ],
    "Labour and delivery": [
        ("Labour care", [
            ("ld_01", "Uses the partograph correctly to monitor labour"),
            ("ld_02", "Conducts a clean, safe delivery with infection prevention practices"),
            ("ld_03", "Provides active management of the third stage of labour"),
            ("ld_04", "Takes the first-line actions for postpartum haemorrhage"),
        ]),
        ("Newborn care and records", [
            ("ld_05", "Provides essential newborn care (drying, warmth, early breastfeeding)"),
            ("ld_06", "Resuscitates a non-breathing newborn with bag and mask"),
            ("ld_07", "Completes the delivery register accurately"),
        ]),
    ],
    "Immunization": [
        ("Vaccine handling", [
            ("imm_01", "Stores and handles vaccines correctly (cold chain, vaccine vial monitor)"),
            ("imm_05", "Disposes of sharps and waste safely"),
        ]),
        ("Vaccination practice", [
            ("imm_02", "Checks the child's immunization status against the schedule"),
            ("imm_03", "Prepares and gives vaccines using the correct technique and site"),
            ("imm_04", "Screens for contraindications and counsels on side effects"),
            ("imm_06", "Records doses in the child card and tally sheets accurately"),
        ]),
    ],
    "HIV care and treatment": [
        ("Testing and counselling", [
            ("hiv_01", "Conducts HIV testing following the national algorithm"),
            ("hiv_02", "Provides pre- and post-test counselling appropriately"),
        ]),
        ("Treatment and follow-up", [
            ("hiv_03", "Assesses clinical stage and starts ART according to guidelines"),
            ("hiv_04", "Monitors viral load and acts on the results"),
            ("hiv_05", "Assesses and supports adherence"),
            ("hiv_06", "Screens for TB and other opportunistic infections"),
            ("hiv_07", "Completes patient files and registers accurately"),
        ]),
    ],
    "TB": [
        ("Case finding and treatment", [
            ("tb_01", "Screens clients for TB symptoms"),
            ("tb_02", "Collects and handles sputum specimens correctly"),
            ("tb_03", "Registers patients and starts treatment per guidelines"),
            ("tb_04", "Monitors adherence and follows up contacts"),
            ("tb_05", "Applies infection prevention and control measures"),
        ]),
    ],
    "Malaria case management": [
        ("Diagnosis and treatment", [
            ("mal_01", "Performs a malaria rapid diagnostic test correctly"),
            ("mal_02", "Interprets results and prescribes the right treatment"),
            ("mal_03", "Recognizes severe malaria and refers"),
            ("mal_04", "Counsels on adherence and prevention"),
        ]),
    ],
    "Data management and reporting": [
        ("Recording and reporting", [
            ("dm_01", "Fills registers completely and consistently"),
            ("dm_02", "Tallies and summarizes data correctly into the monthly report"),
            ("dm_03", "Checks data quality before submitting reports"),
            ("dm_04", "Submits reports on time"),
            ("dm_05", "Uses facility data to review performance and plan"),
        ]),
    ],
}


def default_items() -> pd.DataFrame:
    rows = [(iid, area, domain, text, "yes")
            for area, domains in _SAMPLE.items() for domain, items in domains for iid, text in items]
    return pd.DataFrame(rows, columns=ITEM_COLS)


# ---------------------------------------------------------------------
# Checklist handling
# ---------------------------------------------------------------------
def slug(text) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(text).lower()).strip("_")


def normalize_items(df: pd.DataFrame) -> pd.DataFrame:
    """Trim text, default `active` to yes, drop completely empty rows."""
    out = df.copy()
    for c in ITEM_COLS:
        if c not in out.columns:
            out[c] = ""
    out = out[ITEM_COLS].fillna("").astype(str)
    for c in ITEM_COLS:
        out[c] = out[c].str.strip()
    out = out[~(out[ITEM_COLS] == "").all(axis=1)].reset_index(drop=True)
    out["active"] = out["active"].str.lower().replace({"": "yes", "true": "yes", "false": "no", "1": "yes", "0": "no"})
    return out


def validate_items(df: pd.DataFrame) -> list:
    """Human-readable problems with a checklist (empty list = fine)."""
    errs = []
    d = normalize_items(df)
    if d.empty:
        return ["The checklist is empty."]
    for i, r in d.iterrows():
        line = f"Row {i + 1}"
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", r["item_id"]):
            errs.append(f"{line}: item_id '{r['item_id']}' must start with a letter and use only letters, numbers and underscores.")
        elif r["item_id"].lower() in RESERVED_NAMES or r["item_id"].lower().startswith("grp_"):
            errs.append(f"{line}: item_id '{r['item_id']}' is reserved — choose another.")
        if not r["service_area"]:
            errs.append(f"{line}: service_area is empty.")
        if not r["competency"]:
            errs.append(f"{line}: competency text is empty.")
        if r["active"] not in ("yes", "no"):
            errs.append(f"{line}: active must be yes or no.")
    dup = d["item_id"][d["item_id"].duplicated()].unique()
    if len(dup):
        errs.append("Duplicate item_id: " + ", ".join(dup))
    area_slugs = d.groupby(d["service_area"].map(slug))["service_area"].nunique()
    for s_, n in area_slugs.items():
        if n > 1:
            errs.append(f"Service areas that differ only by punctuation/case clash: '{s_}'.")
    return errs


def load_items(seed: bool = False) -> pd.DataFrame:
    """The checklist from the Sheet. With seed=True an empty checklist is filled with the sample."""
    df = gsheets.get_competency_items()
    if df.empty and seed:
        gsheets.save_competency_items(default_items())
        df = gsheets.get_competency_items()
    return normalize_items(df) if not df.empty else pd.DataFrame(columns=ITEM_COLS)


def active_items(df: pd.DataFrame) -> pd.DataFrame:
    return df[df["active"] != "no"].reset_index(drop=True)


# ---------------------------------------------------------------------
# Saving an in-app assessment
# ---------------------------------------------------------------------
def save_assessment(header: dict, scores: dict, items_df: pd.DataFrame) -> str:
    """header: facility, mentee_name, cadre, service_area, assessment_type (label), assessment_date,
    assessor, visit_ref, comments. scores: {item_id: score code}. Returns the assessment_id."""
    aid = "ca_" + uuid.uuid4().hex[:16]
    text = dict(zip(items_df["item_id"], items_df["competency"]))
    now = datetime.now().isoformat(timespec="seconds")
    rows = []
    for item_id, code in scores.items():
        if code in (None, ""):
            continue
        rows.append({**{k: header.get(k, "") for k in
                        ("visit_ref", "facility", "mentee_name", "cadre", "service_area", "assessment_date",
                         "assessor", "comments")},
                     "assessment_id": aid, "assessment_type": header["assessment_type"],
                     "item_id": item_id, "competency": text.get(item_id, item_id),
                     "score": code, "timestamp": now, "kobo_id": ""})
    gsheets.append_kobo_rows("Competency", rows)
    return aid


# ---------------------------------------------------------------------
# Analysis
# ---------------------------------------------------------------------
def get_assessments() -> pd.DataFrame:
    """All scored items with helper columns: score_num, date_dt, facility_key, mentee_key."""
    df = gsheets.get_all("Competency")
    if df.empty:
        return df
    df = df.drop(columns=["_row"], errors="ignore").copy()
    for c in COMP_COLS:
        if c not in df.columns:
            df[c] = ""
    df["score"] = df["score"].astype(str).str.strip().str.upper()
    df["score_num"] = df["score"].map(SCORE_VALUE)
    df["score_num"] = pd.to_numeric(df["score_num"], errors="coerce")
    df["date_dt"] = pd.to_datetime(df["assessment_date"], errors="coerce")
    df["facility_key"] = df["facility"].astype(str).str.strip().str.lower()
    df["mentee_key"] = df["mentee_name"].astype(str).str.strip().str.lower()
    return df


def summarize(df: pd.DataFrame) -> pd.DataFrame:
    """One row per assessment: % competent, items scored, gaps."""
    cols = ["assessment_id", "visit_ref", "facility", "mentee_name", "cadre", "service_area",
            "assessment_type", "assessment_date", "date_dt", "assessor", "scored", "competent",
            "gaps", "pct_competent"]
    if df.empty:
        return pd.DataFrame(columns=cols)
    d = df.copy()
    d["scored"] = d["score_num"].notna().astype(int)
    d["competent"] = (d["score_num"] >= COMPETENT_MIN).astype(int)
    out = (d.groupby("assessment_id", sort=False)
             .agg(visit_ref=("visit_ref", "first"), facility=("facility", "first"),
                  mentee_name=("mentee_name", "first"), cadre=("cadre", "first"),
                  service_area=("service_area", "first"), assessment_type=("assessment_type", "first"),
                  assessment_date=("assessment_date", "first"), date_dt=("date_dt", "first"),
                  assessor=("assessor", "first"), scored=("scored", "sum"), competent=("competent", "sum"))
             .reset_index())
    out["gaps"] = out["scored"] - out["competent"]
    out["pct_competent"] = (out["competent"] / out["scored"] * 100).where(out["scored"] > 0).round(1)
    return out[cols].sort_values("date_dt", ascending=False, na_position="last").reset_index(drop=True)


def progress(summ: pd.DataFrame) -> pd.DataFrame:
    """Per health worker × service area: first baseline vs latest follow-up."""
    cols = ["facility", "mentee_name", "cadre", "service_area", "baseline_date", "baseline_pct",
            "followup_date", "followup_pct", "change"]
    if summ.empty:
        return pd.DataFrame(columns=cols)
    s = summ.copy()
    s["fkey"] = s["facility"].astype(str).str.strip().str.lower()
    s["mkey"] = s["mentee_name"].astype(str).str.strip().str.lower()
    s = s.sort_values("date_dt", na_position="first")
    keys = ["fkey", "mkey", "service_area"]
    base = s[s["assessment_type"] == TYPE_LABEL["baseline"]].groupby(keys).first()
    foll = s[s["assessment_type"] == TYPE_LABEL["follow_up"]].groupby(keys).last()
    idx = base.index.union(foll.index)
    out = pd.DataFrame(index=idx)
    for tgt, src, col in (("facility", None, "facility"), ("mentee_name", None, "mentee_name"), ("cadre", None, "cadre")):
        out[tgt] = [(base[col].get(i) if i in base.index else foll[col].get(i)) for i in idx]
    out["baseline_date"] = base["assessment_date"].reindex(idx)
    out["baseline_pct"] = base["pct_competent"].reindex(idx)
    out["followup_date"] = foll["assessment_date"].reindex(idx)
    out["followup_pct"] = foll["pct_competent"].reindex(idx)
    out["change"] = (out["followup_pct"] - out["baseline_pct"]).round(1)
    out = out.reset_index()
    return out[["facility", "mentee_name", "cadre", "service_area", "baseline_date", "baseline_pct",
                "followup_date", "followup_pct", "change"]]


def latest_per_worker_area(df: pd.DataFrame) -> pd.DataFrame:
    """Item rows from only the most recent assessment of each health worker in each service area."""
    if df.empty:
        return df
    key = ["facility_key", "mentee_key", "service_area"]
    last = (df.sort_values(["date_dt", "timestamp"], na_position="first")
              .groupby(key).tail(1)[key + ["assessment_id"]].rename(columns={"assessment_id": "_keep"}))
    keep = set(df.merge(last, on=key, how="inner").query("assessment_id == _keep")["assessment_id"])
    return df[df["assessment_id"].isin(keep)]


def weakest(df_latest: pd.DataFrame, area: str, n: int = 10) -> pd.DataFrame:
    d = df_latest[df_latest["service_area"] == area]
    if d.empty:
        return pd.DataFrame(columns=["competency", "assessed", "competent", "pct_competent"])
    g = d.groupby("competency").agg(assessed=("score_num", lambda s: int(s.notna().sum())),
                                    competent=("score_num", lambda s: int((s >= COMPETENT_MIN).sum()))).reset_index()
    g = g[g["assessed"] > 0]
    g["pct_competent"] = (g["competent"] / g["assessed"] * 100).round(1)
    return g.sort_values(["pct_competent", "assessed"], ascending=[True, False]).head(n).reset_index(drop=True)


def gaps_text(scores: dict, items_df: pd.DataFrame) -> str:
    """Bullet list of items scored below competent — the mentorship focus for this visit."""
    text = dict(zip(items_df["item_id"], items_df["competency"]))
    lines = []
    for item_id, code in scores.items():
        v = SCORE_VALUE.get(code)
        if v is not None and v < COMPETENT_MIN:
            lines.append(f"• {text.get(item_id, item_id)} — {SCORE_LABEL[code]}")
    return "\n".join(lines)


# ---------------------------------------------------------------------
# Kobo form (XLSForm) generated from the checklist
# ---------------------------------------------------------------------
def build_xlsform(items_df: pd.DataFrame, form_title: str = "Competency Assessment",
                  form_id: str = "competency_assessment") -> bytes:
    """An .xlsx XLSForm ready to upload to Kobo: facility / health worker details, then one group of
    scored questions per service area that only appears when that area is chosen."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    items = active_items(normalize_items(items_df))
    if items.empty:
        raise ValueError("There are no active checklist items to build a form from.")
    errs = validate_items(items)
    if errs:
        raise ValueError("; ".join(errs[:5]))

    areas = list(dict.fromkeys(items["service_area"]))
    survey = [
        ("text", "facility", "Facility", "Filled in automatically when opened from the dashboard — correct it if needed.", "yes", "", "", ""),
        ("text", "mentee_name", "Health worker being assessed (full name)", "", "yes", "", "", ""),
        ("select_one cadres", "cadre", "Cadre", "", "yes", "", "", ""),
        ("select_one service_areas", "service_area", "Service area assessed", "", "yes", "", "", ""),
        ("select_one assessment_types", "assessment_type", "Assessment type", "Baseline = before mentorship starts.", "yes", "", "", ""),
        ("date", "assessment_date", "Date of assessment", "", "yes", "", "today()", ""),
        ("text", "assessor", "Assessor (your name)", "", "yes", "", "", ""),
        ("text", "visit_ref", "Visit reference (optional)", "Use the same reference for every health worker assessed in one visit.", "", "", "", ""),
    ]
    for area in areas:
        aid = slug(area)
        survey.append(("begin_group", f"grp_{aid}", area, "", "", f"${{service_area}} = '{aid}'", "", ""))
        for _, r in items[items["service_area"] == area].iterrows():
            survey.append(("select_one score_scale", r["item_id"], r["competency"],
                           r["domain"], "yes", "", "", "horizontal-compact"))
        survey.append(("end_group", "", "", "", "", "", "", ""))
    survey.append(("text", "comments", "Comments / observations", "", "", "", "", "multiline"))

    choices = [("cadres", slug(c), c) for c in CADRES]
    choices += [("service_areas", slug(a), a) for a in areas]
    choices += [("assessment_types", c, l) for c, l in ASSESSMENT_TYPES]
    choices += [("score_scale", c, l) for c, l, _ in SCALE]

    wb = Workbook()
    s = wb.active
    s.title = "survey"
    s.append(["type", "name", "label", "hint", "required", "relevant", "default", "appearance"])
    for row in survey:
        s.append(list(row))
    c = wb.create_sheet("choices")
    c.append(["list_name", "name", "label"])
    for row in choices:
        c.append(list(row))
    st_ = wb.create_sheet("settings")
    st_.append(["form_title", "form_id"])
    st_.append([form_title, form_id])
    for ws, widths in ((s, [24, 22, 56, 40, 10, 34, 12, 20]), (c, [18, 30, 44]), (st_, [30, 28])):
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------
# One Kobo submission -> many Sheet rows
# ---------------------------------------------------------------------
def _flat(sub: dict) -> dict:
    """Submission keyed by last path segment, so questions inside groups ('grp_anc/anc_01') match."""
    out = {}
    for k, v in sub.items():
        if isinstance(k, str) and not k.startswith("_"):
            out[k.split("/")[-1]] = v
    return out


def submission_to_rows(sub: dict, items_df: pd.DataFrame, column_map: dict | None = None) -> list:
    """Turn one Kobo competency submission into Sheet rows (one per item answered)."""
    cmap = column_map or {}
    flat = _flat(sub)

    def meta(col, default=""):
        v = flat.get(cmap.get(col, col))
        return default if v is None else str(v).strip()

    kobo_id = str(sub.get("_uuid") or sub.get("_id") or "").strip()
    area_by_slug = {slug(a): a for a in items_df["service_area"].unique()}
    cadre_by_slug = {slug(c): c for c in CADRES}
    area_val = meta("service_area")
    cadre_val = meta("cadre")
    type_val = meta("assessment_type")
    base = {
        "assessment_id": kobo_id,
        "visit_ref": meta("visit_ref"),
        "facility": meta("facility"),
        "mentee_name": meta("mentee_name"),
        "cadre": cadre_by_slug.get(cadre_val, cadre_val),
        "service_area": area_by_slug.get(area_val, area_val),
        "assessment_type": TYPE_LABEL.get(type_val, type_val),
        "assessment_date": meta("assessment_date") or str(sub.get("_submission_time", ""))[:10],
        "assessor": meta("assessor") or str(sub.get("_submitted_by") or ""),
        "comments": meta("comments"),
        "timestamp": str(sub.get("_submission_time", "")),
        "kobo_id": kobo_id,
    }
    rows = []
    for _, it in items_df.iterrows():
        v = flat.get(it["item_id"])
        if v is None or str(v).strip() == "":
            continue
        code = str(v).strip().upper()
        code = {"NA": "NA", "N/A": "NA"}.get(code, str(v).strip())
        rows.append({**base, "item_id": it["item_id"], "competency": it["competency"], "score": code})
    return rows
